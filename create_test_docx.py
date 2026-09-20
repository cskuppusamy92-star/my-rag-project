from pathlib import Path
from docx import Document

output = Path("data/test_documents/test_scheme.docx")

print("Creating:", output.resolve())

document = Document()

document.add_heading("ABC Scheme", level=1)

document.add_paragraph(
    "The ABC Scheme is administered by the Industries Department."
)

document.add_paragraph(
    "Applicants must submit an identity proof and project report."
)

document.add_paragraph(
    "The application should be submitted through the designated portal."
)

document.save(str(output))

print("Created successfully.")
print("Exists:", output.exists())
print("Size:", output.stat().st_size, "bytes")