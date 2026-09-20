from __future__ import annotations

import csv
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from pypdf import PdfReader


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

LIBREOFFICE_PATH = Path(
    os.getenv(
        "LIBREOFFICE_PATH",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
    )
)


# ---------------------------------------------------------------------------
# Common extracted-document model
# ---------------------------------------------------------------------------

@dataclass
class ExtractedDocument:
    """
    Represents one source unit extracted from an uploaded file.

    Every extracted unit keeps its original source location so that
    downstream chunks can retain citation/provenance information.
    """

    text: str
    source_file: str
    source_type: str
    location: dict[str, Any]


# ---------------------------------------------------------------------------
# Utility functions
# ---------------------------------------------------------------------------

def _clean_text(text: str) -> str:
    """Normalize whitespace while preserving readable text."""

    if not text:
        return ""

    lines = []

    for line in text.splitlines():
        cleaned = " ".join(line.split())

        if cleaned:
            lines.append(cleaned)

    return "\n".join(lines)


def _validate_file(path: Path) -> None:
    """Validate that the input file exists and has a supported extension."""

    if not path.exists():
        raise FileNotFoundError(f"File does not exist: {path}")

    if not path.is_file():
        raise ValueError(f"Path is not a file: {path}")

    supported = {".pdf", ".csv",".xlsx", ".txt", ".doc", ".docx"}

    if path.suffix.lower() not in supported:
        raise ValueError(
            f"Unsupported file type: {path.suffix}. "
            f"Supported types: {', '.join(sorted(supported))}"
        )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extract_pdf(path: Path) -> list[ExtractedDocument]:
    """
    Extract text from a PDF page-by-page.

    Each returned item represents one PDF page and contains the
    original page number for citation.
    """

    reader = PdfReader(str(path))

    documents: list[ExtractedDocument] = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = _clean_text(text)

        if not text:
            continue

        documents.append(
            ExtractedDocument(
                text=text,
                source_file=path.name,
                source_type="pdf",
                location={
                    "page": page_number,
                },
            )
        )

    return documents


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def extract_csv(path: Path) -> list[ExtractedDocument]:
    """
    Extract CSV rows while preserving:

    - filename
    - row number
    - column names
    """

    documents: list[ExtractedDocument] = []

    # utf-8-sig handles CSV files exported with a UTF-8 BOM.
    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.DictReader(file)

        if reader.fieldnames is None:
            raise ValueError(f"CSV has no header row: {path.name}")

        columns = [
            column.strip()
            for column in reader.fieldnames
            if column is not None
        ]

        for row_number, row in enumerate(reader, start=2):
            parts = []

            for column in columns:
                value = row.get(column)

                if value is None:
                    continue

                value = str(value).strip()

                if value:
                    parts.append(f"{column}: {value}")

            text = "\n".join(parts)
            text = _clean_text(text)

            if not text:
                continue

            documents.append(
                ExtractedDocument(
                    text=text,
                    source_file=path.name,
                    source_type="csv",
                    location={
                        "row": row_number,
                        "columns": columns,
                    },
                )
            )

    return documents

# ------------------------------------------------------------
# TXT
# ------------------------------------------------------------

def extract_txt(path: Path) -> list[ExtractedDocument]:
    """
    Extract text from a plain text file.
    Each non-empty line becomes one extracted document.
    """

    documents: list[ExtractedDocument] = []

    text = path.read_text(encoding="utf-8")

    for line_number, line in enumerate(text.splitlines(), start=1):
        text_value = _clean_text(line)

        if not text_value:
            continue

        documents.append(
            ExtractedDocument(
                text=text_value,
                source_file=path.name,
                source_type="txt",
                location={
                    "line": line_number
                },
            )
        )

    return documents

# ------------------------------------------------------------
# XLSX
# ------------------------------------------------------------

def extract_xlsx(path: Path) -> list[ExtractedDocument]:
    """
    Extract XLSX rows while preserving:
    - filename
    - sheet name
    - row number
    - column names
    """

    from openpyxl import load_workbook

    documents: list[ExtractedDocument] = []

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )

    for worksheet in workbook.worksheets:

        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            continue

        columns = [
            str(value).strip()
            if value is not None
            else ""
            for value in header_row
        ]

        for row_number, row in enumerate(rows, start=2):

            parts = []

            for column, value in zip(columns, row):

                if not column:
                    continue

                if value is None:
                    continue

                value = str(value).strip()

                if value:
                    parts.append(f"{column}: {value}")

            text = "\n".join(parts)
            text = _clean_text(text)

            if not text:
                continue

            documents.append(
                ExtractedDocument(
                    text=text,
                    source_file=path.name,
                    source_type="xlsx",
                    location={
                        "sheet": worksheet.title,
                        "row": row_number,
                        "columns": columns,
                    },
                )
            )

    workbook.close()

    return documents

# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def extract_docx(path: Path) -> list[ExtractedDocument]:
    """
    Extract DOCX paragraphs while preserving paragraph numbers.
    """

    document = Document(str(path))

    documents: list[ExtractedDocument] = []

    for paragraph_number, paragraph in enumerate(
        document.paragraphs,
        start=1,
    ):
        text = _clean_text(paragraph.text)

        if not text:
            continue

        documents.append(
            ExtractedDocument(
                text=text,
                source_file=path.name,
                source_type="docx",
                location={
                    "paragraph": paragraph_number,
                },
            )
        )

    return documents


# ---------------------------------------------------------------------------
# Legacy DOC
# ---------------------------------------------------------------------------

def _check_libreoffice() -> None:
    """
    Verify that LibreOffice exists before processing a .doc file.
    """

    if not LIBREOFFICE_PATH.exists():
        raise RuntimeError(
            "LibreOffice executable was not found.\n"
            f"Expected location: {LIBREOFFICE_PATH}\n"
            "Install LibreOffice or set LIBREOFFICE_PATH in .env."
        )


def extract_doc(path: Path) -> list[ExtractedDocument]:
    """
    Extract legacy .doc files using LibreOffice.

    We deliberately do NOT treat .doc as .docx.

    LibreOffice converts the legacy document to DOCX in a temporary
    directory. The resulting DOCX is then processed by the dedicated
    DOCX extractor.
    """

    _check_libreoffice()

    with tempfile.TemporaryDirectory(
        prefix="rag_doc_conversion_"
    ) as temp_dir:

        output_dir = Path(temp_dir)

        command = [
            str(LIBREOFFICE_PATH),
            "--headless",
            "--convert-to",
            "docx",
            "--outdir",
            str(output_dir),
            str(path),
        ]

        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )

        if result.returncode != 0:
            raise RuntimeError(
                "LibreOffice failed to convert the .doc file.\n"
                f"stdout: {result.stdout}\n"
                f"stderr: {result.stderr}"
            )

        converted_path = output_dir / f"{path.stem}.docx"

        if not converted_path.exists():
            # LibreOffice may sometimes produce a slightly different
            # filename. Find the generated DOCX as a fallback.
            candidates = list(output_dir.glob("*.docx"))

            if len(candidates) != 1:
                raise RuntimeError(
                    "LibreOffice conversion completed, but the "
                    "converted DOCX could not be located."
                )

            converted_path = candidates[0]

        extracted = extract_docx(converted_path)

        # The source must remain the original .doc file.
        # We therefore rewrite source metadata after conversion.
        for item in extracted:
            item.source_file = path.name
            item.source_type = "doc"

        return extracted


# ---------------------------------------------------------------------------
# Main dispatcher
# ---------------------------------------------------------------------------

def extract_file(path: str | Path) -> list[ExtractedDocument]:
    """
    Extract a supported file using the appropriate parser.

    Supported:

    .pdf
    .csv
    .xlsx
    .txt
    .docx
    .doc
    """

    file_path = Path(path)

    _validate_file(file_path)

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".csv":
        return extract_csv(file_path)

    if extension == ".xlsx":
        return extract_xlsx(file_path)

    if extension == ".txt":
        return extract_txt(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".doc":
        return extract_doc(file_path)

    raise ValueError(
        f"Unsupported file extension: {extension}"
    )


# ---------------------------------------------------------------------------
# Simple summary helper
# ---------------------------------------------------------------------------

def extraction_summary(
    documents: list[ExtractedDocument],
) -> dict[str, Any]:
    """
    Return a simple summary useful for ingestion status messages.
    """

    source_types: dict[str, int] = {}

    total_characters = 0

    for document in documents:
        source_types[document.source_type] = (
            source_types.get(document.source_type, 0) + 1
        )

        total_characters += len(document.text)

    return {
        "units": len(documents),
        "characters": total_characters,
        "source_types": source_types,
    }
